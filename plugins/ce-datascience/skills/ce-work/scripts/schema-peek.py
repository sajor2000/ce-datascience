#!/usr/bin/env python3
"""Print a data file's schema without printing any of its values.

Usage: python3 schema-peek.py <file> [<file> ...] [--delimiter X] [--max-rows N]

Output per file: column heading, inferred type, non-null count, and % missing,
plus the row count. It never prints a cell value, example, min/max, or
distinct value. Small cells are suppressed: a row count under 11 shows as
"<11", and when a column's non-null or missing count is under 11 neither exact
count is shown, so one cannot be recovered from the other by subtraction.

If the first row looks like data (numbers, dates, ID-like or identifier-like
values, duplicate names, or a value repeated in its own column), headings are
printed as col_1..col_n with the reason.

Supports CSV/TSV/PSV (optionally .gz) with the standard library; Parquet,
Feather, and Arrow via pyarrow; and the XLSX header row via openpyxl. Plain
text and other types are refused. Errors print only the exception type.
"""

import csv
import gzip
import os
import re
import sys

# Same list as hooks/data-boundary-guard.py (tested): what counts as a data file.
DATA_EXTS = (
    "csv", "tsv", "psv", "tab", "dat", "parquet", "feather", "arrow", "orc",
    "avro", "xlsx", "xlsm", "xlsb", "xls", "ods", "sas7bdat", "sas7bcat",
    "xpt", "dta", "sav", "zsav", "por", "rds", "rdata", "rda", "fst", "pkl",
    "pickle", "npy", "npz", "h5", "hdf5", "sqlite", "sqlite3", "db", "duckdb",
    "jsonl", "ndjson",
)
THRESHOLD = 11
MISSING = {"", "na", "n/a", "nan", "null", "none", ".", "nat", "#n/a"}
INT_RE = re.compile(r"^[+-]?\d+$")
FLOAT_RE = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
DATE_RE = re.compile(r"^(\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2}(\.\d+)?)?(Z|[+-]\d{2}:?\d{2})?)?"
                     r"|\d{1,2}/\d{1,2}/\d{2,4})$")
BOOL_VALUES = {"true", "false", "t", "f", "yes", "no"}
ID_LIKE = re.compile(r"^[A-Za-z]{0,6}[-_]?\d{3,}$")
IDENTIFIER_LIKE = re.compile(r"@|\d{3}-\d{2}-\d{4}|\(?\d{3}\)?[-. ]\d{3}[-. ]\d{4}")
HEADER_SCAN_ROWS = 1000


def small(n):
    return 0 < n < THRESHOLD


def classify(value):
    if INT_RE.match(value):
        return "integer"
    if FLOAT_RE.match(value):
        return "float"
    if DATE_RE.match(value):
        return "datetime"
    if value.lower() in BOOL_VALUES:
        return "boolean"
    return "string"


def merge_type(current, new):
    if current is None or current == new:
        return new
    if {current, new} == {"integer", "float"}:
        return "float"
    return "string"


def header_problem(names, later_values=None):
    """Return why the first row looks like data, or None when it looks like headings."""
    cleaned = [str(n).strip() if n is not None else "" for n in names]
    for name in cleaned:
        if name and classify(name) in ("integer", "float", "datetime"):
            return "a first-row cell is a number or date"
        if ID_LIKE.match(name):
            return "a first-row cell looks like an ID"
        if IDENTIFIER_LIKE.search(name):
            return "a first-row cell looks like an identifier"
    present = [n.lower() for n in cleaned if n]
    if len(present) != len(set(present)):
        return "the first row has duplicate names"
    if later_values:
        for index, name in enumerate(cleaned):
            if name and name in later_values[index]:
                return "a first-row cell repeats as a value in its column"
    return None


def render(path, rows, columns, note=None):
    """columns: list of (name, type, non_null or None)."""
    print(f"## {os.path.basename(path)}")
    shown_rows = "unknown" if rows is None else (f"<{THRESHOLD}" if small(rows) else str(rows))
    print(f"Rows: {shown_rows}")
    print(f"Small-cell rule: counts from 1 to {THRESHOLD - 1} are shown as <{THRESHOLD}; when a column's "
          "non-null or missing count is that small, neither exact count nor % missing is shown.")
    if note:
        print(note)
    print()
    print("| column | type | non-null | % missing |")
    print("|---|---|---|---|")
    for name, col_type, non_null in columns:
        if non_null == "n/a":
            print(f"| {name} | {col_type} | n/a | n/a |")
            continue
        if non_null is None or rows is None:
            print(f"| {name} | {col_type} | unknown | unknown |")
            continue
        missing = rows - non_null
        if small(non_null) or small(missing):
            shown = f"<{THRESHOLD}" if small(non_null) else "withheld"
            print(f"| {name} | {col_type} | {shown} | withheld |")
            continue
        pct = f"{100.0 * missing / rows:.1f}" if rows else "0.0"
        print(f"| {name} | {col_type} | {non_null} | {pct} |")
    print()


def peek_delimited(path, delimiter, max_rows):
    opener = gzip.open if path.lower().endswith(".gz") else open
    with opener(path, "rt", encoding="utf-8-sig", errors="replace", newline="") as fh:
        reader = csv.reader(fh, delimiter=delimiter)
        first = next(reader, None)
        if first is None:
            render(path, 0, [])
            return
        width = len(first)
        non_null = [0] * width
        types = [None] * width
        seen = [set() for _ in range(width)]
        first_types = [classify(v.strip()) if v.strip().lower() not in MISSING else None for v in first]
        rows = 0
        truncated = False
        for row in reader:
            if max_rows and rows >= max_rows:
                truncated = True
                break
            rows += 1
            for index in range(width):
                value = row[index].strip() if index < len(row) else ""
                if value.lower() in MISSING:
                    continue
                non_null[index] += 1
                types[index] = merge_type(types[index], classify(value))
                if rows <= HEADER_SCAN_ROWS:
                    seen[index].add(value)
    problem = header_problem(first, seen)
    note = None
    if problem:
        note = f"No header detected ({problem}); columns are shown as col_1..col_{width} and the first row is counted as data."
        rows += 1
        for index, value_type in enumerate(first_types):
            if value_type is not None:
                non_null[index] += 1
                types[index] = merge_type(types[index], value_type)
        names = [f"col_{i + 1}" for i in range(width)]
    else:
        names = first
    if truncated:
        note = (note + " " if note else "") + f"Only the first {max_rows} data rows were scanned."
    render(path, rows, [(name, types[i] or "empty", non_null[i]) for i, name in enumerate(names)], note)


def peek_arrow(path, kind):
    try:
        import pyarrow  # noqa: F401
    except ImportError:
        print(f"## {os.path.basename(path)}\npyarrow is not installed; install it in the project "
              "environment to read Parquet/Feather/Arrow schemas.\n")
        return
    if kind == "parquet":
        import pyarrow.parquet as pq
        meta = pq.ParquetFile(path).metadata
        schema = meta.schema.to_arrow_schema()
        rows = meta.num_rows
        leaf = {}
        if meta.num_row_groups:
            group = meta.row_group(0)
            leaf = {group.column(i).path_in_schema: i for i in range(group.num_columns)}
        columns = []
        for field in schema:
            if pyarrow.types.is_nested(field.type) or field.name not in leaf:
                columns.append((field.name, str(field.type), "n/a"))
                continue
            nulls = 0
            for group in range(meta.num_row_groups):
                stats = meta.row_group(group).column(leaf[field.name]).statistics
                if stats is None or not stats.has_null_count:
                    nulls = pq.read_table(path, columns=[field.name]).column(0).null_count
                    break
                nulls += stats.null_count
            columns.append((field.name, str(field.type), rows - nulls))
    else:
        import pyarrow.feather as feather
        table = feather.read_table(path)
        rows = table.num_rows
        columns = [(f.name, str(f.type), rows - table.column(i).null_count)
                   for i, f in enumerate(table.schema)]
    render(path, rows, columns, note="Nested columns report n/a for counts." if any(c[2] == "n/a" for c in columns) else None)


def peek_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        print(f"## {os.path.basename(path)}\nopenpyxl is not installed; install it in the project "
              "environment to read the XLSX header row.\n")
        return
    book = openpyxl.load_workbook(path, read_only=True)
    for sheet in book.worksheets:
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()) or ()
        rows = max((sheet.max_row or 1) - 1, 0)
        typed = [h for h in header if h is not None and not isinstance(h, str)]
        problem = "a first-row cell is a number or date" if typed else header_problem(header)
        if problem:
            names = [f"col_{i + 1}" for i in range(len(header))]
            note = f"XLSX: no header detected ({problem}); approximate row count only."
            rows += 1
        else:
            names = [str(h) if h is not None else "(blank)" for h in header]
            note = "XLSX: header row and approximate row count only."
        render(f"{path} [{sheet.title}]", rows, [(n, "not inferred", None) for n in names], note=note)
    book.close()


def main(argv):
    delimiter, max_rows, paths = None, 0, []
    args = iter(argv)
    for arg in args:
        if arg == "--delimiter":
            delimiter = next(args, ",").encode().decode("unicode_escape")
        elif arg == "--max-rows":
            max_rows = int(next(args, "0") or 0)
        else:
            paths.append(arg)
    if not paths:
        print(__doc__.strip())
        return 2
    status = 0
    for path in paths:
        lower = path.lower()
        base = lower[:-3] if lower.endswith(".gz") else lower
        name = os.path.basename(path)
        try:
            if not os.path.isfile(path):
                print(f"## {name}\nNot a file.\n")
                status = 1
            elif base.endswith((".csv", ".tsv", ".psv", ".tab")):
                default = "\t" if base.endswith((".tsv", ".tab")) else ("|" if base.endswith(".psv") else ",")
                peek_delimited(path, delimiter or default, max_rows)
            elif base.endswith(".parquet"):
                peek_arrow(path, "parquet")
            elif base.endswith((".feather", ".arrow")):
                peek_arrow(path, "feather")
            elif base.endswith((".xlsx", ".xlsm")):
                peek_xlsx(path)
            else:
                ext = base.rsplit(".", 1)[-1] if "." in base else ""
                kind = "a data format schema-peek cannot read" if ext in DATA_EXTS else "not a supported tabular format"
                print(f"## {name}\nRefused: {kind}. Write a script that prints only column names, types, "
                      "row count, and missingness.\n")
                status = 1
        except Exception as exc:  # never echo file bytes from an error message
            print(f"## {name}\nCould not read the file ({type(exc).__name__}).\n")
            status = 1
    return status


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
